# -*- coding: utf-8 -*-
"""新疆科技学院本科毕业论文（设计）构建脚本 —— Markdown 分章草稿 → 符合 2026 版撰写规范的 Word 成品。

用法（在论文项目根目录运行）：
    python build_thesis.py [--src thesis] [--out thesis/build/论文.docx] [--refresh-toc] [--pdf]

依赖：python-docx、latex2mathml、lxml；公式转换需本机安装 Word（使用其 MML2OMML.XSL）。
配置：<src>/thesis.json（封面信息）。首次运行若不存在会生成配置模板后退出。

排版规则实现依据《新疆科技学院本科毕业论文（设计）撰写规范（2026版修订）》：
  A4；上下 2.5cm 左右 3.0cm；页眉页脚 1.5cm；宋体 + Times New Roman；行距多倍 1.3。
  章标题小二粗体居中(段前后0.5行)；节标题小三粗体顶格；三级四号粗体；四级小四。
  摘要/目录/参考文献/致谢/附录标题小三粗体居中；正文小四首行缩进2字符。
  图题在图下、表题在表上，五号居中，按章编号；表为三线表；公式右端编号。
  页眉：正文奇数页=论文题目、偶数页=校名；摘要目录无页眉；目录罗马页码、正文起阿拉伯页码。
"""
import argparse
import copy
import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---------- 字号（磅） ----------
SZ_XIAOER = 18      # 小二
SZ_XIAOSAN = 15     # 小三
SZ_SIHAO = 14       # 四号
SZ_XIAOSI = 12      # 小四
SZ_WUHAO = 10.5     # 五号

CONTENT_WIDTH_CM = 15.0   # A4(21cm) - 左右各3cm
ASSETS = Path(__file__).resolve().parent.parent / "assets"

WRITING_DECLARATION = (
    "本人郑重声明：所呈交的学位论文，是本人在导师的指导下，独立进行研究工作所取得的成果。"
    "除文中已经注明引用的内容外，本论文不含任何其他个人或集体已经发表或撰写过的作品或成果。"
    "对本文的研究做出重要贡献的个人和集体，均已在文中以明确方式标明。本声明的法律结果由本人承担。"
)
LICENSE_DECLARATION = (
    "本人完全了解新疆科技学院关于收集、保存、使用学位论文的规定，即：按照学校要求提交学位论文的印刷本和电子版本；"
    "学校有权保存学位论文的印刷本和电子版，并提供目录检索与阅览服务；学校可以采用影印、缩印、数字化或其它复制手段保存论文；"
    "在不以赢利为目的的前提下，学校可以将学位论文编入有关数据库，提供网上服务。（保密论文在解密后遵守此规定）"
)

# ---------- 公式：LaTeX -> OMML ----------
_MML_XSL_CANDIDATES = [
    r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files\Microsoft Office\Office16\MML2OMML.XSL",
    r"C:\Program Files\Microsoft Office\root\Office15\MML2OMML.XSL",
]
_omml_transform = None
_omml_warned = False


def _get_omml_transform():
    global _omml_transform
    if _omml_transform is not None:
        return _omml_transform
    try:
        from lxml import etree
        xsl = next((p for p in _MML_XSL_CANDIDATES if Path(p).exists()), None)
        if xsl is None:
            return None
        _omml_transform = etree.XSLT(etree.parse(xsl))
    except Exception:
        return None
    return _omml_transform


def latex_to_omml(latex):
    """返回可直接挂到 w:p 下的 m:oMath 元素；失败返回 None。"""
    global _omml_warned
    transform = _get_omml_transform()
    if transform is None:
        if not _omml_warned:
            print("[警告] 未找到 Word 的 MML2OMML.XSL，公式将以斜体文本代替（请安装 Word 或手工在 Word 中重录公式）")
            _omml_warned = True
        return None
    try:
        from lxml import etree
        import latex2mathml.converter
        mathml = latex2mathml.converter.convert(latex)
        omml_doc = transform(etree.fromstring(mathml))
        node = omml_doc.getroot()
        # 转成 python-docx 可用的 lxml 元素副本
        return copy.deepcopy(node)
    except Exception as exc:
        print(f"[警告] 公式转换失败，以斜体文本代替: {latex[:50]} ({exc})")
        return None


# ---------- 底层格式助手 ----------

def set_run_font(run, size=SZ_XIAOSI, bold=False, italic=False, cjk="宋体", ascii_font="Times New Roman"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = ascii_font
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cjk)


def _spacing_el(p):
    ppr = p._p.get_or_add_pPr()
    sp = ppr.find(qn("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing")
        ppr.append(sp)
    return sp


def fmt_para(p, align=None, line13=True, before_lines=None, after_lines=None,
             first_indent_chars=None, hanging_chars=None, keep_with_next=False,
             page_break_before=False):
    """按规范设置段落：行距多倍1.3、段前后（单位0.01行）、首行/悬挂缩进（单位0.01字符）。"""
    if align is not None:
        p.alignment = align
    sp = _spacing_el(p)
    if line13:
        sp.set(qn("w:line"), "312")          # 1.3 * 240
        sp.set(qn("w:lineRule"), "auto")
    if before_lines is not None:
        sp.set(qn("w:beforeLines"), str(before_lines))
        sp.set(qn("w:before"), "0")
    if after_lines is not None:
        sp.set(qn("w:afterLines"), str(after_lines))
        sp.set(qn("w:after"), "0")
    if first_indent_chars is not None or hanging_chars is not None:
        ppr = p._p.get_or_add_pPr()
        ind = ppr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            ppr.append(ind)
        if first_indent_chars is not None:
            ind.set(qn("w:firstLineChars"), str(first_indent_chars))
        if hanging_chars is not None:
            ind.set(qn("w:hangingChars"), str(hanging_chars))
            ind.set(qn("w:hanging"), str(int(hanging_chars * 0.01 * SZ_WUHAO * 20)))
    if keep_with_next:
        p.paragraph_format.keep_with_next = True
    if page_break_before:
        p.paragraph_format.page_break_before = True


def set_outline_level(p, level=0):
    """让非 Heading 样式的标题（摘要/参考文献等）也进入目录。"""
    ppr = p._p.get_or_add_pPr()
    el = OxmlElement("w:outlineLvl")
    el.set(qn("w:val"), str(level))
    ppr.append(el)


def add_field(p, instr, hint_text=None, size=SZ_XIAOSI):
    """插入 Word 域（PAGE / TOC 等）。"""
    r1 = p.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    fld.set(qn("w:dirty"), "true")
    r1._element.append(fld)
    r2 = p.add_run()
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    r2._element.append(instr_el)
    r3 = p.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3._element.append(sep)
    if hint_text:
        rt = p.add_run(hint_text)
        set_run_font(rt, size=size)
    r4 = p.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r4._element.append(end)


def setup_page(sec):
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(3.0)
    sec.header_distance = Cm(1.5)
    sec.footer_distance = Cm(1.5)


def clear_part(part):
    part.is_linked_to_previous = False
    if not part.paragraphs:
        part.add_paragraph()
    first = part.paragraphs[0]
    for p in list(part.paragraphs):
        if p is not first:
            p._p.getparent().remove(p._p)
    first.text = ""


def blank_headers_footers(sec):
    for part in (sec.header, sec.even_page_header, sec.footer, sec.even_page_footer):
        clear_part(part)


def _first_para(part):
    if not part.paragraphs:
        part.add_paragraph()
    return part.paragraphs[0]


def set_header_text(part, text):
    clear_part(part)
    p = _first_para(part)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, size=SZ_WUHAO)


def set_footer_pagenum(part):
    clear_part(part)
    p = _first_para(part)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, " PAGE ", hint_text="1", size=SZ_WUHAO)
    for r in p.runs:
        r.font.size = Pt(SZ_WUHAO)


def set_pgnum(sec, fmt=None, start=None):
    sectpr = sec._sectPr
    el = sectpr.find(qn("w:pgNumType"))
    if el is None:
        el = OxmlElement("w:pgNumType")
        sectpr.append(el)
    if fmt:
        el.set(qn("w:fmt"), fmt)
    if start is not None:
        el.set(qn("w:start"), str(start))


# ---------- 行内 Markdown（**粗体** `代码` $公式$） ----------

_INLINE_SPLIT = re.compile(r"(\$[^$]+\$|\*\*[^*]+?\*\*|`[^`]+`)")


def add_inline(p, text, size=SZ_XIAOSI, bold=False):
    for part in _INLINE_SPLIT.split(text):
        if not part:
            continue
        if part.startswith("$") and part.endswith("$") and len(part) > 2:
            omml = latex_to_omml(part[1:-1])
            if omml is not None:
                p._p.append(omml)
            else:
                r = p.add_run(part[1:-1])
                set_run_font(r, size=size, italic=True)
        elif part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            set_run_font(r, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            set_run_font(r, size=size, ascii_font="Consolas")
        else:
            r = p.add_run(part)
            set_run_font(r, size=size, bold=bold)


# ---------- Markdown 块解析 ----------

def _join_lines(lines):
    """中文行拼接不加空格，西文行拼接加空格。"""
    out = ""
    for ln in lines:
        ln = ln.strip()
        if not out:
            out = ln
        elif out and (_is_cjk(out[-1]) or _is_cjk(ln[0])):
            out += ln
        else:
            out += " " + ln
    return out


def _is_cjk(ch):
    return "一" <= ch <= "鿿" or ch in "，。；：、（）《》“”—…！？"


def parse_blocks(text):
    """把一个章节 Markdown 解析为块列表。

    块类型：("h", level, text) ("p", text) ("table", rows) ("figure", caption, path)
            ("math", latex, tag) ("code", text) ("item", text)
    以 > 开头的引用行是作者备注，不进入论文。
    """
    blocks = []
    lines = text.splitlines()
    i = 0
    para_buf = []

    def flush():
        if para_buf:
            blocks.append(("p", _join_lines(para_buf)))
            para_buf.clear()

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s or s.startswith(">") or s.startswith("<!--"):
            flush()
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            flush()
            blocks.append(("h", len(m.group(1)), m.group(2).strip()))
            i += 1
            continue
        if s.startswith("```"):
            flush()
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", "\n".join(code)))
            continue
        if s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", s)
        if m:
            flush()
            blocks.append(("figure", m.group(1).strip(), m.group(2).strip()))
            i += 1
            continue
        if s.startswith("$$"):
            flush()
            body = s[2:]
            math_lines = []
            if body.rstrip().endswith("$$"):
                math_lines.append(body.rstrip()[:-2])
                i += 1
            else:
                if body:
                    math_lines.append(body)
                i += 1
                while i < len(lines) and not lines[i].strip().endswith("$$"):
                    math_lines.append(lines[i])
                    i += 1
                if i < len(lines):
                    math_lines.append(lines[i].strip()[:-2])
                    i += 1
            latex = " ".join(x.strip() for x in math_lines).strip()
            tag = None
            mt = re.search(r"\\tag\{([^}]*)\}", latex)
            if mt:
                tag = mt.group(1)
                latex = latex.replace(mt.group(0), "").strip()
            blocks.append(("math", latex, tag))
            continue
        m = re.match(r"^([-*]|\d+[.)])\s+(.*)$", s)
        if m:
            flush()
            marker = m.group(1)
            item_text = m.group(2)
            if re.match(r"\d", marker):
                item_text = f"（{marker.rstrip('.)')}）{item_text}"
            blocks.append(("item", item_text))
            i += 1
            continue
        para_buf.append(s)
        i += 1
    flush()
    return blocks


def normalize_heading(text):
    """规范：编号后空两格书写标题名称。"""
    return re.sub(r"^(\d+(?:\.\d+)*)\s+", r"\1  ", text)


# ---------- 内容渲染 ----------

def add_body_para(doc, text, size=SZ_XIAOSI, indent=True):
    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             first_indent_chars=200 if indent else None)
    add_inline(p, text, size=size)
    return p


def add_special_heading(doc, text, before=100, after=200, size=SZ_XIAOSAN, outline=0):
    """摘要/目录/参考文献/致谢/附录式标题：小三粗体居中。"""
    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, before_lines=before, after_lines=after,
             keep_with_next=True)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=True)
    if outline is not None:
        set_outline_level(p, outline)
    return p


def add_chapter_heading(doc, text, page_break=False):
    p = doc.add_paragraph(style="Heading 1")
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, before_lines=50, after_lines=50,
             keep_with_next=True, page_break_before=page_break)
    r = p.add_run(normalize_heading(text))
    set_run_font(r, size=SZ_XIAOER, bold=True)
    return p


def add_sub_heading(doc, level, text):
    style = {2: "Heading 2", 3: "Heading 3", 4: "Heading 4"}[level]
    size = {2: SZ_XIAOSAN, 3: SZ_SIHAO, 4: SZ_XIAOSI}[level]
    bold = level != 4
    p = doc.add_paragraph(style=style)
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, before_lines=0, after_lines=0,
             keep_with_next=True)
    r = p.add_run(normalize_heading(text))
    set_run_font(r, size=size, bold=bold)
    return p


def add_caption(doc, text, above=True):
    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=above,
             before_lines=25 if above else 0, after_lines=0 if above else 25)
    # 规范：图/表序后空两格书写题名
    m = re.match(r"^(图|表)\s*(\d+\.\d+)\s*(.*)$", text)
    if m:
        text = f"{m.group(1)}{m.group(2)}  {m.group(3)}"
    r = p.add_run(text)
    set_run_font(r, size=SZ_WUHAO)
    return p


def _set_cell_borders(cell, top=None, bottom=None):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for name, sz in (("top", top), ("bottom", bottom)):
        if sz is None:
            continue
        el = borders.find(qn(f"w:{name}"))
        if el is None:
            el = OxmlElement(f"w:{name}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), "000000")


def add_three_line_table(doc, rows):
    """三线表：顶线/底线 1.5 磅，栏目线 0.75 磅；内容五号居中；表头跨页重复。"""
    ncols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=ncols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = None
    tblpr = tbl._tbl.tblPr
    # 无整体边框（逐单元格画三线）；自动调整列宽
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "autofit")
    tblpr.append(layout)
    tbl.autofit = True
    for ri, row in enumerate(rows):
        if ri == 0:
            trpr = tbl.rows[0]._tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader")
            trpr.append(th)
        for ci in range(ncols):
            cell = tbl.cell(ri, ci)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = row[ci] if ci < len(row) else ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # 表格内为默认单倍行距（规范）
            add_inline(p, text, size=SZ_WUHAO, bold=(ri == 0))
            _set_cell_borders(cell,
                              top=12 if ri == 0 else None,
                              bottom=12 if ri == len(rows) - 1 else (6 if ri == 0 else None))
    return tbl


def add_figure(doc, caption, img_path, base_dir):
    path = Path(img_path)
    if not path.is_absolute():
        path = (base_dir / img_path).resolve()
    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)
    if path.exists():
        run = p.add_run()
        pic = run.add_picture(str(path))
        if pic.width > Cm(14):
            ratio = Cm(14) / pic.width
            pic.width = Cm(14)
            pic.height = int(pic.height * ratio)
    else:
        r = p.add_run(f"【缺图：{img_path}】")
        set_run_font(r, size=SZ_WUHAO, bold=True)
        print(f"[警告] 图片不存在: {img_path}")
    if caption:
        add_caption(doc, caption, above=False)


def add_display_math(doc, latex, tag):
    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent_chars=500)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    omml = latex_to_omml(latex)
    if omml is not None:
        p._p.append(omml)
    else:
        r = p.add_run(latex)
        set_run_font(r, italic=True)
    if tag:
        r = p.add_run(f"\t（{tag}）")
        set_run_font(r, size=SZ_XIAOSI)


def add_code_block(doc, text):
    for line in text.splitlines():
        p = doc.add_paragraph()
        fmt_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, line13=False)
        r = p.add_run(line if line.strip() else " ")
        set_run_font(r, size=SZ_WUHAO, ascii_font="Consolas")


def render_blocks(doc, blocks, base_dir, first_h1_page_break):
    """渲染一个章节文件的块。表题(前置段落"表x.y …")自动与表绑定。"""
    seen_h1 = False
    i = 0
    while i < len(blocks):
        blk = blocks[i]
        kind = blk[0]
        if kind == "h":
            level, text = blk[1], blk[2]
            if level == 1:
                add_chapter_heading(doc, text, page_break=first_h1_page_break or seen_h1)
                seen_h1 = True
            else:
                add_sub_heading(doc, min(level, 4), text)
        elif kind == "p":
            nxt = blocks[i + 1] if i + 1 < len(blocks) else None
            if nxt and nxt[0] == "table" and re.match(r"^表\s*\d+\.\d+", blk[1]):
                add_caption(doc, blk[1], above=True)
                add_three_line_table(doc, nxt[1])
                i += 2
                continue
            add_body_para(doc, blk[1])
        elif kind == "table":
            add_three_line_table(doc, blk[1])
        elif kind == "figure":
            add_figure(doc, blk[1], blk[2], base_dir)
        elif kind == "math":
            add_display_math(doc, blk[1], blk[2])
        elif kind == "code":
            add_code_block(doc, blk[1])
        elif kind == "item":
            add_body_para(doc, blk[1])
        i += 1


# ---------- 各大部件 ----------

def build_cover(doc, cfg):
    sec = doc.sections[0]
    sec.start_type = WD_SECTION_START.NEW_PAGE
    setup_page(sec)
    blank_headers_footers(sec)
    for _ in range(2):
        p = doc.add_paragraph()
        fmt_para(p)
    logo = ASSETS / "cover-logo.png"
    if logo.exists():
        p = doc.add_paragraph()
        fmt_para(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        p.add_run().add_picture(str(logo), width=Cm(12.5))
    for _ in range(2):
        fmt_para(doc.add_paragraph())
    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("本科毕业论文（设计）")
    set_run_font(r, size=26, bold=True)
    for _ in range(3):
        fmt_para(doc.add_paragraph())
    fields = [("题    目", cfg.get("title", "")),
              ("学    院", cfg.get("college", "")),
              ("专业班级", cfg.get("major_class", "")),
              ("姓    名", cfg.get("student_name", "")),
              ("学    号", cfg.get("student_id", "")),
              ("指导教师", cfg.get("advisor", ""))]
    tbl = doc.add_table(rows=len(fields), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = None
    for ri, (label, value) in enumerate(fields):
        lc, vc = tbl.cell(ri, 0), tbl.cell(ri, 1)
        lc.width = Cm(3.6)
        vc.width = Cm(8.5)
        lp = lc.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = lp.add_run(f"{label}：")
        set_run_font(r, size=SZ_XIAOSAN, bold=True)
        vp = vc.paragraphs[0]
        vp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = vp.add_run(value)
        set_run_font(r, size=SZ_XIAOSAN)
        _set_cell_borders(vc, bottom=8)
        for cell in (lc, vc):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            sp = _spacing_el(cell.paragraphs[0])
            sp.set(qn("w:line"), "440")
            sp.set(qn("w:lineRule"), "atLeast")
    for _ in range(3):
        fmt_para(doc.add_paragraph())
    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(cfg.get("date", ""))
    set_run_font(r, size=SZ_SIHAO, bold=True)


def build_declaration(doc):
    sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    setup_page(sec)
    blank_headers_footers(sec)
    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, before_lines=50, after_lines=100)
    r = p.add_run("新疆科技学院学位论文写作声明和使用授权说明")
    set_run_font(r, size=16, bold=True)

    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, before_lines=50, after_lines=50)
    r = p.add_run("学位论文写作声明")
    set_run_font(r, size=SZ_SIHAO, bold=True)
    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent_chars=200)
    r = p.add_run(WRITING_DECLARATION)
    set_run_font(r, size=SZ_XIAOSI)
    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.RIGHT, before_lines=100, after_lines=200)
    r = p.add_run("论文作者签名：          日期：    年    月    日")
    set_run_font(r, size=SZ_XIAOSI)

    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, before_lines=100, after_lines=50)
    r = p.add_run("学位论文使用授权说明")
    set_run_font(r, size=SZ_SIHAO, bold=True)
    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent_chars=200)
    r = p.add_run(LICENSE_DECLARATION)
    set_run_font(r, size=SZ_XIAOSI)
    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, before_lines=100)
    r = p.add_run("论文作者签名：            导师签名：")
    set_run_font(r, size=SZ_XIAOSI)
    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.RIGHT, before_lines=50)
    r = p.add_run("日期：    年    月    日")
    set_run_font(r, size=SZ_XIAOSI)


def _abstract_sections(md_text):
    """从摘要 md 中取出 (中文段落列表, 中文关键词, 英文段落列表, 英文关键词)。"""
    cn, kw_cn, en, kw_en = [], "", [], ""
    mode = None
    for blk in parse_blocks(md_text):
        if blk[0] == "h":
            t = blk[2].replace(" ", "").lower()
            if "摘要" in t and blk[1] >= 2:
                mode = "cn"
            elif "abstract" in t:
                mode = "en"
            continue
        if blk[0] not in ("p", "item"):
            continue
        text = blk[1]
        m = re.match(r"^\*?\*?关键词\*?\*?\s*[:：]\s*(.*)$", text)
        if m:
            kw_cn = m.group(1).strip()
            continue
        m = re.match(r"^\*?\*?key\s*words?\*?\*?\s*[:：]\s*(.*)$", text, re.I)
        if m:
            kw_en = m.group(1).strip()
            continue
        if mode == "cn":
            cn.append(text)
        elif mode == "en":
            en.append(text)
    return cn, kw_cn, en, kw_en


def build_abstract(doc, cfg, md_text):
    cn, kw_cn, en, kw_en = _abstract_sections(md_text)
    # 中文摘要
    sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    setup_page(sec)
    blank_headers_footers(sec)
    set_pgnum(sec, fmt="upperRoman", start=1)  # 摘要不印页码，但让目录中显示为 I/II
    p = add_special_heading(doc, "摘  要", before=50, after=50)
    for text in cn:
        p = doc.add_paragraph()
        fmt_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent_chars=200,
                 before_lines=50, after_lines=50)
        add_inline(p, text, size=SZ_XIAOSI)
    fmt_para(doc.add_paragraph())  # 空一行
    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, before_lines=50, after_lines=50)
    r = p.add_run("关键词：")
    set_run_font(r, size=SZ_XIAOSI, bold=True)
    r = p.add_run(kw_cn)
    set_run_font(r, size=SZ_XIAOSI)
    # 英文摘要
    sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    setup_page(sec)
    blank_headers_footers(sec)
    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, before_lines=50, after_lines=50)
    r = p.add_run(cfg.get("title_en", ""))
    set_run_font(r, size=SZ_XIAOSAN)
    p = add_special_heading(doc, "ABSTRACT", before=50, after=50)
    for text in en:
        p = doc.add_paragraph()
        fmt_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent_chars=200,
                 before_lines=50, after_lines=50)
        add_inline(p, text, size=SZ_XIAOSI)
    fmt_para(doc.add_paragraph())
    p = doc.add_paragraph()
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, before_lines=50, after_lines=50)
    r = p.add_run("Keywords：")
    set_run_font(r, size=SZ_XIAOSI, bold=True)
    r = p.add_run(kw_en)
    set_run_font(r, size=SZ_XIAOSI)


def build_toc(doc):
    sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    setup_page(sec)
    blank_headers_footers(sec)
    set_footer_pagenum(sec.footer)
    set_footer_pagenum(sec.even_page_footer)
    set_pgnum(sec, fmt="upperRoman", start=1)
    add_special_heading(doc, "目  录", before=100, after=200, outline=None)
    p = doc.add_paragraph()
    fmt_para(p)
    add_field(p, ' TOC \\o "1-3" \\h \\z \\u ',
              hint_text="（目录为自动域：在 Word 中按 Ctrl+A 后按 F9 更新）")


def build_body(doc, cfg, chapter_texts, src_dir):
    sec = doc.add_section(WD_SECTION_START.ODD_PAGE)
    setup_page(sec)
    doc_type = cfg.get("doc_type", "论文")
    even_text = "新疆科技学院本科毕业论文" if doc_type == "论文" else "新疆科技学院本科毕业设计"
    set_header_text(sec.header, cfg.get("title", ""))          # 奇数页：论文题目
    set_header_text(sec.even_page_header, even_text)            # 偶数页：校名
    set_footer_pagenum(sec.footer)
    set_footer_pagenum(sec.even_page_footer)
    set_pgnum(sec, fmt="decimal", start=1)
    for idx, text in enumerate(chapter_texts):
        render_blocks(doc, parse_blocks(text), src_dir, first_h1_page_break=(idx > 0))


def build_tail_section(doc, header_text, continue_footer=True):
    sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    setup_page(sec)
    set_header_text(sec.header, header_text)
    set_header_text(sec.even_page_header, header_text)
    if continue_footer:
        set_footer_pagenum(sec.footer)
        set_footer_pagenum(sec.even_page_footer)
    return sec


def build_conclusion(doc, md_text, src_dir):
    build_tail_section(doc, "结  论")
    blocks = parse_blocks(md_text)
    p = doc.add_paragraph(style="Heading 1")
    fmt_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, before_lines=50, after_lines=50)
    r = p.add_run("结  论")
    set_run_font(r, size=SZ_XIAOER, bold=True)
    for blk in blocks:
        if blk[0] == "h" and blk[1] == 1:
            continue  # 标题已手工渲染
        render_blocks(doc, [blk], src_dir, first_h1_page_break=False)


def build_references(doc, md_text):
    build_tail_section(doc, "参考文献")
    add_special_heading(doc, "参考文献", before=100, after=200)
    entries = []
    for ln in md_text.splitlines():
        s = ln.strip()
        if re.match(r"^\[\d+\]", s):
            entries.append(s.replace("★", "").replace("  ", " ").strip())
    for entry in entries:
        p = doc.add_paragraph()
        fmt_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, hanging_chars=280)
        r = p.add_run(entry)
        set_run_font(r, size=SZ_WUHAO)
    return len(entries)


def build_acknowledgement(doc, md_text, src_dir):
    build_tail_section(doc, "致  谢")
    add_special_heading(doc, "致  谢", before=100, after=200)
    for blk in parse_blocks(md_text):
        if blk[0] == "h":
            continue
        render_blocks(doc, [blk], src_dir, first_h1_page_break=False)


def build_appendix(doc, md_text, src_dir):
    build_tail_section(doc, "附  录")
    add_special_heading(doc, "附  录", before=100, after=200)
    for blk in parse_blocks(md_text):
        if blk[0] == "h" and blk[1] == 1:
            continue
        render_blocks(doc, [blk], src_dir, first_h1_page_break=False)


# ---------- 全局样式 ----------

def setup_styles(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(SZ_XIAOSI)
    st.element.get_or_add_rPr()
    rfonts = st.element.rPr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        st.element.rPr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "宋体")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    for name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        hs = doc.styles[name]
        hs.font.color.rgb = RGBColor(0, 0, 0)
        hs.font.name = "Times New Roman"
        rpr = hs.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        rf.set(qn("w:eastAsia"), "宋体")
    # 奇偶页眉不同 + 打开时提示更新域（刷新目录）
    settings = doc.settings.element
    if settings.find(qn("w:evenAndOddHeaders")) is None:
        settings.append(OxmlElement("w:evenAndOddHeaders"))
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)


# ---------- 素材发现与主流程 ----------

def classify_md(path):
    text = path.read_text(encoding="utf-8")
    for ln in text.splitlines():
        s = ln.strip()
        if s.startswith("#"):
            head = s.lstrip("#").strip()
            if head.startswith("摘要"):
                return "abstract", text
            if head.startswith("参考文献"):
                return "references", text
            if head.startswith("致谢") or head.startswith("致 谢"):
                return "ack", text
            if head.startswith("结论") or head.startswith("结 论"):
                return "conclusion", text
            if head.startswith("附录") or head.startswith("附 录"):
                return "appendix", text
            if re.match(r"^\d+\s", head):
                return "chapter", text
            return "skip", text
    return "skip", text


DEFAULT_CONFIG = {
    "title": "论文题目（≤25字）",
    "title_en": "English Title (no more than 12 notional words)",
    "doc_type": "论文",
    "college": "学院名称",
    "major_class": "专业班级",
    "student_name": "姓名",
    "student_id": "学号",
    "advisor": "指导教师",
    "date": "二〇二六年五月",
    "output": "thesis/build/毕业论文.docx",
}


def refresh_fields_via_word(docx_path, export_pdf=False):
    """用本机 Word 更新目录/页码域；可选导出 PDF。"""
    pdf_path = str(Path(docx_path).with_suffix(".pdf"))
    ps = f'''
$word = New-Object -ComObject Word.Application
$word.Visible = $false
try {{
  $doc = $word.Documents.Open("{docx_path}")
  foreach ($toc in $doc.TablesOfContents) {{ $toc.Update() | Out-Null }}
  $doc.Fields.Update() | Out-Null
  $doc.Repaginate()
  foreach ($toc in $doc.TablesOfContents) {{ $toc.UpdatePageNumbers() | Out-Null }}
  $doc.Save()
  {f'$doc.SaveAs2("{pdf_path}", 17)' if export_pdf else ''}
  $doc.Close($false)
  Write-Output "fields refreshed"
}} finally {{ $word.Quit() }}
'''
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps],
                       capture_output=True, text=True)
    print(r.stdout.strip() or r.stderr.strip())


def main():
    ap = argparse.ArgumentParser(description="Markdown → 新疆科技学院毕业论文 Word")
    ap.add_argument("--src", default="thesis", help="Markdown 章节所在目录")
    ap.add_argument("--out", default=None, help="输出 docx 路径（默认取配置 output）")
    ap.add_argument("--refresh-toc", action="store_true", help="生成后用 Word 更新目录页码域")
    ap.add_argument("--pdf", action="store_true", help="同时导出 PDF（隐含 --refresh-toc）")
    args = ap.parse_args()

    src = Path(args.src)
    cfg_path = src / "thesis.json"
    if not cfg_path.exists():
        cfg_path.write_text(json.dumps(DEFAULT_CONFIG, ensure_ascii=False, indent=2),
                            encoding="utf-8")
        print(f"[!] 未找到配置，已生成模板 {cfg_path}，请填写封面信息后重新运行。")
        sys.exit(1)
    cfg = json.loads(cfg_path.read_text(encoding="utf-8"))

    parts = {"chapters": []}
    for md in sorted(src.glob("[0-9][0-9]_*.md")):
        if md.name.startswith("00"):
            continue
        kind, text = classify_md(md)
        if kind == "chapter":
            parts["chapters"].append(text)
        elif kind != "skip":
            parts[kind] = text
        print(f"  {md.name} -> {kind}")

    missing = [k for k in ("abstract", "references", "ack") if k not in parts]
    if missing or not parts["chapters"]:
        print(f"[!] 缺少必要部分: {missing or ''} 章节数={len(parts['chapters'])}")
        if not parts["chapters"]:
            sys.exit(1)

    doc = Document()
    setup_styles(doc)
    build_cover(doc, cfg)
    build_declaration(doc)
    if "abstract" in parts:
        build_abstract(doc, cfg, parts["abstract"])
    build_toc(doc)
    build_body(doc, cfg, parts["chapters"], src)
    if "conclusion" in parts:
        build_conclusion(doc, parts["conclusion"], src)
    n_refs = build_references(doc, parts["references"]) if "references" in parts else 0
    if "ack" in parts:
        build_acknowledgement(doc, parts["ack"], src)
    if "appendix" in parts:
        build_appendix(doc, parts["appendix"], src)

    out = Path(args.out or cfg.get("output", "thesis/build/毕业论文.docx"))
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"[OK] 已生成 {out}（章节 {len(parts['chapters'])}，参考文献 {n_refs} 条）")
    print("     目录/页码为 Word 域，打开文档时会提示更新；或使用 --refresh-toc 自动更新。")

    if args.refresh_toc or args.pdf:
        refresh_fields_via_word(str(out.resolve()), export_pdf=args.pdf)
        if args.pdf:
            print(f"[OK] PDF: {out.with_suffix('.pdf')}")


if __name__ == "__main__":
    main()
