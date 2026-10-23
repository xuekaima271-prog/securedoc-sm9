# -*- coding: utf-8 -*-
"""按学校官方开题报告表生成马学凯的开题报告 docx。
底表：1.新疆科技学院...开题报告（学生填写）.docx
规范：信息安全毕业论文开题报告规范2025.01.23（小五、1.3倍行距、首行缩进、悬挂缩进参考文献）
"""
import shutil, re
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"D:\微信\xwechat_files\wxid_tgzypa3l8qkj22_e28f\msg\file\2026-09\1.新疆科技学院毕业论文（设计）毕业论文（设计）开题报告（学生填写）.docx"
OUT = r"D:\biyesheji\开题报告\2311131081+马学凯_开题报告.docx"

BODY_SIZE = 9  # 小五

def fmt_run(r, size=BODY_SIZE, bold=False):
    r.font.name = 'Times New Roman'
    r.font.size = Pt(size)
    r.font.bold = bold
    rPr = r._element.get_or_add_rPr()
    rPr.get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')

def set_para(p, first_line_chars=None, hanging_chars=None, line=1.3,
             align=None, size=BODY_SIZE, bold=False):
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if align is not None:
        pf.alignment = align
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    if first_line_chars:
        ind.set(qn('w:firstLine'), str(int(first_line_chars * size * 20)))
        ind.set(qn('w:firstLineChars'), str(first_line_chars * 100))
    if hanging_chars:
        ind.set(qn('w:hanging'), str(int(hanging_chars * size * 20)))
        ind.set(qn('w:hangingChars'), str(hanging_chars * 100))

def add_para(cell, text, first_line_chars=None, hanging_chars=None,
             bold=False, align=None, size=BODY_SIZE):
    p = cell.add_paragraph()
    set_para(p, first_line_chars=first_line_chars, hanging_chars=hanging_chars,
             align=align, size=size, bold=bold)
    r = p.add_run(text)
    fmt_run(r, size=size, bold=bold)
    return p

def clear_cell(cell):
    """remove all paragraphs except the first, then empty it"""
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p0 = cell.paragraphs[0]
    for r in list(p0.runs):
        r._element.getparent().remove(r._element)
    return p0

def set_value_cell(cell, text, bold=False, size=12):
    p0 = clear_cell(cell)
    set_para(p0, line=1.3, align=WD_ALIGN_PARAGRAPH.CENTER, size=size, bold=bold)
    r = p0.add_run(text)
    fmt_run(r, size=size, bold=bold)

# ---------------- 内容 ----------------
TITLE = "基于国密SM9的端到端加密文档投递系统"
COLLEGE = "信息科学与工程学院"
CLASS_ = "信息安全2023-2班"
NAME = "马学凯"
STUID = "2311131081"

S1 = [
("h", "一、选题背景及研究意义"),
("p", "随着企业数字化转型的深入，合同文本、财务报表、研发图纸、人事档案等敏感文档越来越多地以电子形式在组织内部流转。文档在“发送—存储—接收”全链路上面临三类典型威胁：传输链路被窃听导致内容外泄，内部人员越权访问不属于自己的文档，以及存储服务器失陷或运维人员滥用权限造成批量泄露。近年来国内外多起数据泄露事件表明，来自服务端与内部人员的威胁危害往往大于外部攻击，仅依靠边界防护难以根治[1]。"),
("p", "目前企业常用的文档保护方案是“传输层加密+服务端存储加密”，其局限在于加解密密钥通常由服务器掌握，文档可能在服务端被还原，数据库管理员、运维人员或取得服务端权限的攻击者因而具有接触明文的技术条件。端到端加密（End-to-End Encryption，E2EE）将内容加解密移至通信端点，服务端只负责密文传递与存储，可显著缩小服务端明文暴露面[2]。"),
("p", "与此同时，《中华人民共和国密码法》自2020年1月1日起施行[3]，商用密码应用安全性评估逐步落地，GB/T 39786—2021对信息系统密码应用提出了明确的技术要求[4]。SM3杂凑、SM4分组与SM9标识密码算法均已发布为国家标准[5][6]，其中SM9以用户身份标识直接参与公钥计算，无需证书签发与分发，天然契合“以员工身份为中心”的企业加密文档投递场景。"),
("p", "本课题设计并实现一个基于国密SM9的端到端加密文档投递系统，其意义在于：其一，验证真实SM9标识密码运算在浏览器环境中运行的工程可行性，使文档加解密真正发生在用户终端；其二，给出“业务数据库与对象存储不持久化明文、文件密钥和用户身份私钥”的盲存储设计；其三，为中小企业提供一条不依赖专用密码硬件的国密文档安全投递技术路线。"),
("h", "二、国内外研究现状"),
("h2", "（一）国内研究现状"),
("p", "端到端加密与安全投递方面，国内企业网盘与协同办公产品普遍采用服务端加密模式，服务提供方持有密钥；部分安全厂商提供基于SM2/SM4的文档加密网关，但加解密发生在网关侧而非用户终端，本质上仍未实现端到端加密。"),
("p", "标识密码方面，袁峰、程朝辉对SM9标识密码算法的设计原理、安全基础与应用场景进行了系统综述，为后续工程实现提供了完整参考[7]。赖建昌等人提出了基于商密SM9的高效标识广播加密方案，将SM9的密钥封装能力扩展到多方接收场景[8]。冯登国等人对云计算安全的研究指出，数据加密与密钥管理是云端数据保护的核心问题，服务端持钥模式存在固有风险[2]。在标准层面，国家密码管理局及全国信安标委发布的GM/T 0044—2016与GB/T 38635系列标准明确了SM9算法机制[5][6]，程朝辉进一步从英文文献角度系统整理了SM9密码方案的设计与安全基础[15]。"),
("h2", "（二）国外研究现状"),
("p", "标识密码（Identity-Based Cryptography，IBC）思想最早由Shamir于1984年提出，其核心是让用户身份标识直接参与公钥确定，从而减少证书签发与分发开销[9]。Boneh与Franklin利用椭圆曲线上的Weil配对构造出首个实用的标识加密方案（BF-IBE），奠定了基于双线性对的标识密码技术路线[10]；Barreto与Naehrig提出的配对友好曲线进一步提升了配对运算效率，成为高效实现的基础[11]。"),
("p", "端到端加密方面，Cohn-Gordon等人对Signal协议进行了形式化安全分析，验证了其核心安全属性[12]；Unger等人系统梳理了主流安全消息协议的设计空间与安全目标[13]；Ermoshina等人综述了端到端加密消息协议的技术脉络，指出端到端保护正从即时消息向文件共享等场景扩展[14]。"),
("p", "综合来看，现有工作仍存在不足：主流企业方案密钥集中于服务端，服务器可接触明文；已有端到端加密产品多采用国际算法体系，基于SM9的端到端文档投递缺少公开的完整工程实现；基于PKI的方案需为每位员工签发证书，管理成本高；此外多数系统的操作日志可被具有数据库权限的管理员无痕篡改，审计可信度不足。本课题即针对上述问题展开研究。"),
("h", "三、研究的基本内容与拟解决的主要问题"),
("h2", "（一）基本内容"),
("p", "本课题设计并实现基于国密SM9的企业级端到端加密文档投递系统，采用“浏览器加密端—盲存储应用服务器—密钥生成中心（KGC）”分域架构，文档内容加解密与密钥封装全部在浏览器内完成。主要研究内容包括："),
("p", "1.浏览器内国密混合加密模块：基于SM4-CBC分块加密文件、SM3校验密文完整性、SM9-IBE按接收人身份封装文件密钥，支持“一份密文、多人接收”，加解密经WebWorker后台线程执行。"),
("p", "2.KGC密钥中心：主私钥经PBKDF2派生后以AES-256-GCM加密落盘，按身份确定性提取用户私钥，支持换设备重签发与签发审计。"),
("p", "3.盲存储应用服务：业务数据库与对象存储仅持久化密文、加密元数据与密钥信封，支持分块上传下载与本地/MinIO存储切换。"),
("p", "4.安全机制与功能实现：SM3审计哈希链、方法级访问控制、七类安全事件记录，以及用户、文档、审计与系统管理功能。"),
("p", "5.测试验证：完成单元与集成测试、端到端场景测试、国密算法标准符合性验证与性能测试。"),
("h2", "（二）拟解决的主要问题"),
("p", "1.浏览器端缺少成熟可用的SM9密码库，双线性对运算开销较大，如何在浏览器中实现正确且性能可接受的真实SM9标识加解密并支持大文件处理，是本课题需要首先解决的问题。"),
("p", "2.多接收人场景下，如何在服务器不接触明文、文件密钥与用户身份私钥的前提下，完成文件密钥的安全分发、密文存储与元数据管理，是本课题的核心问题。"),
("p", "3.普通数据库日志可被管理员无痕篡改，用户换设备后面临私钥恢复困难，如何保障审计记录的抗抵赖性与密钥的可恢复性，需要在设计中专门解决。"),
("h", "四、拟采用的研究方法及预期成果"),
("h2", "（一）研究方法"),
("p", "1.文献研究法：系统查阅国密算法标准、标识密码与端到端加密相关文献，梳理国内外研究现状，确定技术路线与研究方案。"),
("p", "2.系统分析法：从功能需求、非功能需求与安全威胁三个维度对系统进行分析，明确角色用例与“服务器不接触明文”等安全红线。"),
("p", "3.原型开发法：采用Spring Boot与Vue3技术栈迭代开发，先实现加密核心模块与最小业务闭环，再逐步完善管理、审计与存储功能。"),
("p", "4.实验验证法：构建单元、集成与端到端三层测试体系，对SM3/SM4/SM9关键运算进行标准向量对标，并通过基准测试评估加密与传输性能。"),
("h2", "（二）预期成果"),
("p", "1.一套功能完整的基于国密SM9的端到端加密文档投递系统，涵盖浏览器加密端、盲存储应用服务器与KGC密钥中心，实现密钥初始化、加密投递、接收解密、多人接收、审计追溯等完整业务闭环。"),
("p", "2.测试报告一份：单元与集成测试全部通过，端到端场景全部检查点通过，并给出SM3、SM4、SM9算法标准符合性与性能测试数据。"),
("p", "3.符合规范要求的毕业论文一篇，系统总结研究背景、方案设计、实现过程与测试结论。"),
("h", "五、论文（设计）计划安排"),
("p", "2026.12.23—2026.12.31 根据课题查阅相关文献，撰写开题报告，开题报告定稿。"),
("p", "2027.01.01—2027.03.09 进行课题关键技术的研究、实验设计、系统开发、实验结果的处理与分析，完成毕业论文（设计）初稿。"),
("p", "2027.03.10—2027.03.14 准备毕业论文（设计）中期检查，填写过程性材料，并修改论文。"),
("p", "2027.03.15—2027.04.25 修改并完成毕业论文（设计）终稿。"),
("p", "2027.04.26—2027.05.12 提交毕业论文（设计），做好答辩PPT，完成毕业论文（设计）答辩。"),
("p", "2027.05.13—2027.05.30 完成毕业论文（设计）相关材料的整理及归档工作。"),
("h", "六、主要参考文献"),
("r", "[1] 冯登国, 张敏, 李昊. 大数据安全与隐私保护[J]. 计算机学报, 2014, 37(1): 246-258."),
("r", "[2] 冯登国, 张敏, 张妍, 等. 云计算安全研究[J]. 软件学报, 2011, 22(1): 71-83."),
("r", "[3] 全国人民代表大会常务委员会. 中华人民共和国密码法[Z]. 北京: 人民出版社, 2019."),
("r", "[4] 国家市场监督管理总局, 国家标准化管理委员会. GB/T 39786—2021 信息安全技术 信息系统密码应用基本要求[S]. 北京: 中国标准出版社, 2021."),
("r", "[5] 国家市场监督管理总局, 国家标准化管理委员会. GB/T 38635.2—2020 信息安全技术 SM9标识密码算法 第2部分: 算法[S]. 北京: 中国标准出版社, 2020."),
("r", "[6] 国家密码管理局. GM/T 0044—2016 SM9标识密码算法[S]. 北京: 中国标准出版社, 2016."),
("r", "[7] 袁峰, 程朝辉. SM9标识密码算法综述[J]. 信息安全研究, 2016, 2(11): 1008-1027."),
("r", "[8] 赖建昌, 黄欣沂, 何德彪. 一种基于商密SM9的高效标识广播加密方案[J]. 计算机学报, 2021, 44(5): 897-907."),
("r", "[9] SHAMIR A. Identity-based cryptosystems and signature schemes[C]//Advances in Cryptology: CRYPTO 1984. Berlin: Springer, 1985: 47-53."),
("r", "[10] BONEH D, FRANKLIN M. Identity-based encryption from the Weil pairing[C]//Advances in Cryptology: CRYPTO 2001. Berlin: Springer, 2001: 213-229."),
("r", "[11] BARRETO P S L M, NAEHRIG M. Pairing-friendly elliptic curves of prime order[C]//Selected Areas in Cryptography: SAC 2005. Berlin: Springer, 2006: 319-331."),
("r", "[12] COHN-GORDON K, CREMERS C, DOWLING B, et al. A formal security analysis of the Signal messaging protocol[C]//2017 IEEE European Symposium on Security and Privacy (EuroS&P). Paris: IEEE, 2017: 451-466."),
("r", "[13] UNGER N, DECHAND S, BONNEAU J, et al. SoK: Secure Messaging[C]//2015 IEEE Symposium on Security and Privacy. San Jose: IEEE, 2015: 232-249."),
("r", "[14] ERMOSHINA K, MUSIANI F, HALPIN H. End-to-end encrypted messaging protocols: an overview[C]//Internet Science: INSCI 2016. Cham: Springer, 2016: 244-254."),
("r", "[15] CHENG Z. The SM9 cryptographic schemes[EB/OL]. (2017-02-13). https://eprint.iacr.org/2017/117."),
]

# ---------------- 生成 ----------------
import os
if os.path.exists(OUT):
    os.chmod(OUT, 0o666)
    os.remove(OUT)
doc = docx.Document(SRC)
t = doc.tables[0]

set_value_cell(t.rows[0].cells[2], TITLE, bold=True, size=12)   # 题目
set_value_cell(t.rows[1].cells[2], COLLEGE, size=12)             # 二级学院
set_value_cell(t.rows[1].cells[4], CLASS_, size=12)              # 专业班级
set_value_cell(t.rows[2].cells[2], NAME, size=12)                # 学生姓名
set_value_cell(t.rows[2].cells[4], STUID, size=12)               # 学号
# 指导教师 / 职称 留空，由用户填写

body = t.rows[4].cells[0]
clear_cell(body)

for kind, text in S1:
    if kind == "h":
        add_para(body, text, bold=True)
    elif kind == "h2":
        add_para(body, text, bold=True)
    elif kind == "r":
        add_para(body, text, hanging_chars=2)
    else:
        add_para(body, text, first_line_chars=2)

doc.save(OUT)
print("saved:", OUT)

# 字数统计（正文，不含参考文献）
cn_chars = 0
for kind, text in S1:
    if kind in ("p", "h", "h2"):
        cn_chars += len(re.findall(r'[\u4e00-\u9fff]', text))
print("正文字数（汉字，不含参考文献）:", cn_chars)
total = sum(len(re.findall(r'[\u4e00-\u9fff]', text)) for _, text in S1)
print("全文汉字数:", total)
