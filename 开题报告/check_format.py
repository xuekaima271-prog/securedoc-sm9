# -*- coding: utf-8 -*-
"""Check formatting of the generated docx against the spec."""
import docx
from docx.oxml.ns import qn

path = r"D:\biyesheji\开题报告\2311131081+马学凯_开题报告.docx"
doc = docx.Document(path)
t = doc.tables[0]

def run_info(p):
    out = []
    for r in p.runs:
        rPr = r._element.find(qn('w:rPr'))
        if rPr is None:
            out.append(("?", None, None))
            continue
        rf = rPr.find(qn('w:rFonts'))
        ascii_f = rf.get(qn('w:ascii')) if rf is not None else None
        ea = rf.get(qn('w:eastAsia')) if rf is not None else None
        sz = rPr.find(qn('w:sz'))
        szv = sz.get(qn('w:val')) if sz is not None else None
        # w:b 存在不等于加粗：python-docx 设 bold=False 会写 <w:b w:val="0"/>
        b = rPr.find(qn('w:b'))
        if b is None:
            bold = False
        else:
            bold = b.get(qn('w:val')) not in ("0", "false")
        out.append((ascii_f, ea, szv, bold))
    return out

def para_info(p):
    pPr = p._p.find(qn('w:pPr'))
    ind = pPr.find(qn('w:ind')) if pPr is not None else None
    fl = ind.get(qn('w:firstLineChars')) if ind is not None else None
    hg = ind.get(qn('w:hangingChars')) if ind is not None else None
    spacing = pPr.find(qn('w:spacing')) if pPr is not None else None
    line = spacing.get(qn('w:line')) if spacing is not None else None
    lineRule = spacing.get(qn('w:lineRule')) if spacing is not None else None
    return fl, hg, line, lineRule

body = t.rows[4].cells[0]
print("body paragraphs:", len(body.paragraphs))
issues = []
for i, p in enumerate(body.paragraphs):
    txt = p.text.strip()
    if not txt:
        continue
    fl, hg, line, lineRule = para_info(p)
    if line is None or line != "312":  # 1.3 * 240
        issues.append(f"para{i} line spacing wrong: {line} {lineRule}")
    ri = run_info(p)
    for j, (a, e, s, b) in enumerate(ri):
        if a != 'Times New Roman':
            issues.append(f"para{i} run{j} ascii font: {a}")
        if e != '宋体':
            issues.append(f"para{i} run{j} eastAsia font: {e}")
        if s != '18':  # 9pt = 18 half-points
            issues.append(f"para{i} run{j} size: {s}")
    # heading vs body indent check
    if txt.startswith(("一、","二、","三、","四、","五、","六、")) and txt.endswith(("研究意义","研究现状","主要问题","预期成果","计划安排","参考文献")):
        pass
    print(f"[{i}] fl={fl} hg={hg} line={line}/{lineRule} sz={ri[0][2] if ri else '?'} bold={ri[0][3] if ri else '?'} | {txt[:36]}")

# header value cells
for ri, ci in [(0,2),(1,2),(1,4),(2,2),(2,4)]:
    c = t.rows[ri].cells[ci]
    p = c.paragraphs[0]
    ri_ = run_info(p)
    al = p.paragraph_format.alignment
    print(f"cell({ri},{ci}) '{p.text[:20]}' fmt={ri_} align={al}")

print()
if issues:
    print("ISSUES:")
    for x in issues:
        print(" -", x)
else:
    print("no font/spacing issues found")
