# -*- coding: utf-8 -*-
"""Inspect the school form table structure (merges, fonts) before filling."""
import docx
from docx.oxml.ns import qn

path = r"D:\微信\xwechat_files\wxid_tgzypa3l8qkj22_e28f\msg\file\2026-09\1.新疆科技学院毕业论文（设计）毕业论文（设计）开题报告（学生填写）.docx"
doc = docx.Document(path)
t = doc.tables[0]
print("rows:", len(t.rows), "cols:", len(t.columns))
for ri, row in enumerate(t.rows):
    cells = row.cells
    seen = {}
    for ci, c in enumerate(cells):
        tc = c._tc
        tid = id(tc)
        if tid not in seen:
            seen[tid] = {"first_col": ci, "text": c.text.strip()[:40].replace("\n", "|")}
        else:
            seen[tid]["last_col"] = ci
    desc = "; ".join(f"col{info['first_col']}-{info.get('last_col', info['first_col'])}:{info['text']}" for info in seen.values())
    print(f"row{ri}: {desc}")
    for ci, c in enumerate(row.cells):
        tcPr = c._tc.tcPr
        gs = tcPr.find(qn('w:gridSpan')) if tcPr is not None else None
        vm = tcPr.find(qn('w:vMerge')) if tcPr is not None else None
        if gs is not None or vm is not None:
            gsv = gs.get(qn('w:val')) if gs is not None else None
            vmv = vm.get(qn('w:val')) if vm is not None else 'continue'
            print(f"  col{ci}: gridSpan={gsv} vMerge={vmv}")

# paragraph count in body cell (row4 merged)
print("\nbody cell paragraphs:", len(t.rows[4].cells[1].paragraphs))

# inspect fonts of label cell and default style
def show_run_fmt(p, label):
    for r in p.runs[:2]:
        rPr = r._element.find(qn('w:rPr'))
        fonts = []
        if rPr is not None:
            rf = rPr.find(qn('w:rFonts'))
            if rf is not None:
                fonts = [rf.get(qn('w:ascii')), rf.get(qn('w:eastAsia')), rf.get(qn('w:hAnsi'))]
            sz = rPr.find(qn('w:sz'))
            szv = sz.get(qn('w:val')) if sz is not None else None
        print(f"  {label}: text='{r.text[:20]}' fonts={fonts} sz={szv}")

p0 = doc.paragraphs[0]
print("\ndoc title para runs:")
show_run_fmt(p0, "title")
p1 = t.rows[0].cells[0].paragraphs[0]
print("table row0 label runs:")
show_run_fmt(p1, "label")
p_body = t.rows[4].cells[1].paragraphs[0]
print("body cell first para runs:")
show_run_fmt(p_body, "body")

# section properties: page size / margins
sec = doc.sections[0]
from docx.shared import Cm
print("\npage w/h:", sec.page_width.cm, sec.page_height.cm, "margins L/R/T/B:", sec.left_margin.cm, sec.right_margin.cm, sec.top_margin.cm, sec.bottom_margin.cm)
